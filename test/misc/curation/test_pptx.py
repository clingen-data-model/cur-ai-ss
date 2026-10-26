import io

from PIL import Image
from pptx import Presentation

from lib.misc.curation.models import CurationSummaryRow, SectionContent
from lib.misc.curation.pptx import build_curation_pptx

PUBLICATION = [SectionContent(title='Publication', content='Doe et al., 2024')]


def _row(patient_id: int, proband: str, variant: str) -> CurationSummaryRow:
    return CurationSummaryRow(
        paper_id=1,
        patient_id=patient_id,
        publication_and_testing=PUBLICATION,
        proband=[SectionContent(title='', content=proband)],
        variant_info=[SectionContent(title=variant, content='Zygosity: Het')],
        clinical_presentation=[SectionContent(title='', content='Seizures')],
        functional_segregation=[],
        score=[SectionContent(title='Family "A"', content='1.0')],
    )


def _tables(pptx_bytes: bytes) -> list:
    prs = Presentation(io.BytesIO(pptx_bytes))
    return [
        shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table
    ]


def test_one_slide_per_proband_each_with_headers_and_publication() -> None:
    rows = [
        _row(10, 'P1', 'c.1A>G'),
        _row(10, 'P1', 'c.2A>G'),
        _row(11, 'P2', 'c.3A>G'),
    ]

    tables = _tables(build_curation_pptx(rows))

    assert len(tables) == 2
    for table in tables:
        assert [cell.text for cell in table.rows[0].cells][:2] == [
            'Publication & Testing',
            'Proband',
        ]
        assert 'Doe et al., 2024' in table.cell(1, 0).text
    assert len(tables[0].rows) == 3  # header + one row per variant
    assert len(tables[1].rows) == 2
    assert 'c.1A>G' in tables[0].cell(1, 2).text
    assert 'c.2A>G' in tables[0].cell(2, 2).text
    assert 'c.3A>G' in tables[1].cell(1, 2).text


def test_proband_columns_merge_down_the_variant_rows() -> None:
    rows = [_row(10, 'P1', 'c.1A>G'), _row(10, 'P1', 'c.2A>G')]

    (table,) = _tables(build_curation_pptx(rows))

    for col in (0, 1, 3, 5):
        assert table.cell(1, col).is_merge_origin
        assert table.cell(2, col).is_spanned
    for col in (2, 4):
        assert not table.cell(1, col).is_merge_origin
    assert table.cell(1, 1).text.count('P1') == 1


def test_pedigree_goes_on_every_slide(tmp_path) -> None:
    image = tmp_path / 'pedigree.png'
    Image.new('RGB', (200, 600)).save(image)
    rows = [_row(10, 'P1', 'c.1A>G'), _row(11, 'P2', 'c.2A>G')]
    for row in rows:
        row.pedigree_image_path = str(image)

    prs = Presentation(io.BytesIO(build_curation_pptx(rows)))

    for slide in prs.slides:
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        assert len(pictures) == 1
        assert pictures[0].height <= prs.slide_height - pictures[0].top


def test_no_rows_gives_an_empty_deck() -> None:
    assert len(Presentation(io.BytesIO(build_curation_pptx([]))).slides) == 0
